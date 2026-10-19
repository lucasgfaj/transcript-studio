"""Renders the minutes Markdown as a Word document (headings, lists, checkboxes, bold, tables)."""
from __future__ import annotations

import io
import re

from docx import Document
from docx.shared import Pt
from docx.text.paragraph import Paragraph

_BOLD = re.compile(r"(\*\*[^*]+\*\*|__[^_]+__)")
_CHECKBOX = re.compile(r"^\[( |x|X)\]\s+")
_LIST = re.compile(r"^(\s*)([-*+]|\d+[.)])\s+(.*)$")


def _add_inline(paragraph: Paragraph, text: str) -> None:
    for part in _BOLD.split(text):
        if not part:
            continue
        if _BOLD.fullmatch(part):
            paragraph.add_run(part[2:-2]).bold = True
        else:
            paragraph.add_run(part.replace("`", ""))


def _table(document: Document, rows: list[str]) -> None:
    cells = [[cell.strip() for cell in row.strip().strip("|").split("|")] for row in rows
             if not re.fullmatch(r"\s*\|?[\s:|-]+\|?\s*", row)]
    if not cells:
        return
    table = document.add_table(rows=len(cells), cols=max(len(r) for r in cells))
    table.style = "Table Grid"
    for i, row in enumerate(cells):
        for j, value in enumerate(row):
            paragraph = table.cell(i, j).paragraphs[0]
            _add_inline(paragraph, value)
            if i == 0:
                for run in paragraph.runs:
                    run.bold = True


class DocxRenderer:
    def render(self, markdown: str) -> bytes:
        document = Document()
        document.styles["Normal"].font.name = "Calibri"
        document.styles["Normal"].font.size = Pt(11)
        lines = markdown.splitlines()
        i = 0
        while i < len(lines):
            line = lines[i].rstrip()
            if line.lstrip().startswith("|"):
                block = []
                while i < len(lines) and lines[i].lstrip().startswith("|"):
                    block.append(lines[i])
                    i += 1
                _table(document, block)
                continue
            i += 1
            if not line.strip() or line.strip() in ("---", "***"):
                continue
            if heading := re.match(r"^(#{1,6})\s+(.*)$", line):
                document.add_heading(heading[2].strip("*").strip(), level=min(len(heading[1]), 4))
                continue
            if item := _LIST.match(line):
                text = item[3]
                level = len(item[1].replace("\t", "  ")) // 2
                style = "List Number" if item[2][0].isdigit() else "List Bullet"
                if box := _CHECKBOX.match(text):
                    text = ("☑ " if box[1].lower() == "x" else "☐ ") + text[box.end():]
                    style = "List Paragraph"
                if level and style != "List Paragraph":
                    style = f"{style} 2"
                _add_inline(document.add_paragraph(style=style), text)
                continue
            if line.startswith(">"):
                _add_inline(document.add_paragraph(style="Quote"), line.lstrip("> "))
                continue
            _add_inline(document.add_paragraph(), line.strip())
        buffer = io.BytesIO()
        document.save(buffer)
        return buffer.getvalue()
