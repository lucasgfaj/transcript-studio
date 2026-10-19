from __future__ import annotations

import io
import json
import threading
import time
from typing import Any

from docx import Document
from fastapi.testclient import TestClient

from transcript_studio.core.config import Settings
from transcript_studio.main import create_app
from transcript_studio.services.minutes_writer import (
    MINUTES_TOKENS, NOTE_TOKENS, MinutesRequest, MinutesWriter, split_text,
)
from tests.conftest import FakeChat, FakeGpu, FakeTranscriber, make_wav


def transcribed_meeting(client: TestClient, text: str = "Bruno: o orçamento está aprovado.") -> dict[str, Any]:
    meeting = client.post("/api/meetings", json={"title": "Obra", "participants": ["Ana", "Bruno"],
                                                 "date": "2026-10-01"}).json()
    client.put(f"/api/meetings/{meeting['id']}/transcript", json={"text": text})
    return meeting


def wait_for(client: TestClient, job_id: str, status: str, timeout: float = 5) -> dict[str, Any]:
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        job = client.get(f"/api/jobs/{job_id}").json()
        if job["status"] == status:
            return job
        time.sleep(0.02)
    raise AssertionError(f"job never reached {status}: {job}")


def test_generates_minutes_as_a_new_version_and_frees_whisper(client: TestClient, chat: FakeChat,
                                                              transcriber: FakeTranscriber) -> None:
    meeting = transcribed_meeting(client)
    client.put(f"/api/meetings/{meeting['id']}/minutes", json={"markdown": "# Rascunho manual"})

    chat.gate = threading.Event()
    response = client.post(f"/api/meetings/{meeting['id']}/minutes/generate", json={"model": "qwen3.5:4b"})
    assert response.status_code == 202
    assert client.get(f"/api/meetings/{meeting['id']}").json()["status"] == "generating"
    chat.gate.set()
    job = wait_for(client, response.json()["id"], "done")

    assert "qwen3.5:4b" in job["message"]
    assert transcriber.unloads >= 1
    model, system, prompt = chat.calls[-1]
    assert model == "qwen3.5:4b" and "português do Brasil" in system
    assert "o orçamento está aprovado" in prompt and "Participantes: Ana, Bruno" in prompt

    bundle = client.get(f"/api/meetings/{meeting['id']}/minutes").json()
    assert bundle["current"]["markdown"].startswith("# Ata") and bundle["current"]["model"] == "qwen3.5:4b"
    assert len(bundle["history"]) == 2
    old = client.get(f"/api/meetings/{meeting['id']}/minutes/{bundle['history'][1]['id']}").json()
    assert old["markdown"] == "# Rascunho manual\n"
    assert client.get(f"/api/meetings/{meeting['id']}").json()["status"] == "done"


def test_generation_requires_a_transcript(client: TestClient) -> None:
    meeting = client.post("/api/meetings", json={"title": "Vazia"}).json()
    assert client.post(f"/api/meetings/{meeting['id']}/minutes/generate").status_code == 409


def test_stream_sends_tokens_and_late_subscribers_get_text_so_far(client: TestClient, chat: FakeChat) -> None:
    chat.gate = threading.Event()
    meeting = transcribed_meeting(client)
    job = client.post(f"/api/meetings/{meeting['id']}/minutes/generate").json()
    deadline = time.monotonic() + 5
    while not chat.calls and time.monotonic() < deadline:
        time.sleep(0.01)
    time.sleep(0.3)  # let the first token batch flush

    events: list[dict[str, Any]] = []
    with client.stream("GET", f"/api/jobs/{job['id']}/events") as stream:
        for line in stream.iter_lines():
            if not line.startswith("data: "):
                continue
            events.append(json.loads(line[6:]))
            if events[-1]["type"] == "snapshot":
                chat.gate.set()
            if events[-1]["type"] == "done":
                break

    text = events[0]["text"] + "".join(e["text"] for e in events if e["type"] == "token")
    assert events[0]["text"].startswith("#")  # what was streamed before this subscriber connected
    assert text == chat.answer
    assert events[-1]["job"]["status"] == "done"


def test_cancel_generation_keeps_previous_status(client: TestClient, chat: FakeChat) -> None:
    chat.gate = threading.Event()
    meeting = transcribed_meeting(client)
    job = client.post(f"/api/meetings/{meeting['id']}/minutes/generate").json()
    while not chat.calls:
        time.sleep(0.01)
    client.post(f"/api/jobs/{job['id']}/cancel")
    wait_for(client, job["id"], "cancelled")
    assert client.get(f"/api/meetings/{meeting['id']}").json()["status"] == "transcribed"
    assert client.get(f"/api/meetings/{meeting['id']}/minutes").json()["current"] is None


def test_long_transcripts_use_map_reduce() -> None:
    chat = FakeChat()
    transcript = "\n\n".join(f"Parágrafo {i}: " + "texto " * 40 for i in range(6))
    writer = MinutesWriter(chat, chunk_characters=600)
    steps: list[str] = []
    markdown = writer.write(MinutesRequest("Obra", (), "2026-10-01", transcript), "m", on_step=lambda m, _p: steps.append(m))
    chunks = split_text(transcript, 600)
    assert len(chunks) > 1
    assert len(chat.calls) == len(chunks) + 1
    assert steps[0] == f"Lendo trecho 1 de {len(chunks)}" and steps[-1] == "Escrevendo a ata"
    assert "NOTAS PARCIAIS" in chat.calls[-1][2]
    assert chat.max_tokens[0] == NOTE_TOKENS and chat.max_tokens[-1] == MINUTES_TOKENS
    assert markdown.startswith("# Ata")


class VerboseChat(FakeChat):
    """A model whose condensed notes never get shorter must not loop forever."""

    def stream(self, model, system, prompt, should_stop=None, max_tokens=None):  # type: ignore[no-untyped-def]
        self.calls.append((model, system, prompt))
        if "Elabore a ata" in prompt:
            yield "# Ata"
        else:
            yield "nota longa " * 400


def test_condensing_stops_when_notes_do_not_shrink() -> None:
    chat = VerboseChat()
    transcript = "\n\n".join("texto " * 100 for _ in range(6))
    writer = MinutesWriter(chat, chunk_characters=600, synthesis_characters=1_000)
    assert writer.write(MinutesRequest("Obra", (), "2026-10-01", transcript), "m").startswith("# Ata")
    condense_calls = [c for c in chat.calls if "Una e condense" in c[2]]
    notes = "\n\n--- NOTA PARCIAL ---\n\n".join(["nota longa " * 400] * len(split_text(transcript, 600)))
    # Exactly one round: its output did not shrink, so a second round is never attempted.
    assert len(condense_calls) == len(split_text(notes, 1_000))


def test_exports(client: TestClient) -> None:
    meeting = client.post("/api/meetings/from-audio", files={"file": ("obra.wav", make_wav(1), "audio/wav")},
                          data={"title": "Obra Sul", "date": "2026-10-01"}).json()
    job = client.post(f"/api/meetings/{meeting['id']}/transcribe").json()
    wait_for(client, job["id"], "done")
    client.put(f"/api/meetings/{meeting['id']}/minutes",
               json={"markdown": "# Ata\n\n## Tarefas\n\n- [x] Feito **já**\n- [ ] Pendente\n\n| A | B |\n|---|---|\n| 1 | 2 |"})

    md = client.get(f"/api/meetings/{meeting['id']}/export", params={"format": "md"})
    assert md.status_code == 200 and "obra-sul-2026-10-01-ata.md" in md.headers["content-disposition"]

    docx = client.get(f"/api/meetings/{meeting['id']}/export", params={"format": "docx"})
    document = Document(io.BytesIO(docx.content))
    texts = [p.text for p in document.paragraphs]
    assert "Ata" in texts and "☑ Feito já" in texts and "☐ Pendente" in texts
    assert document.tables[0].cell(1, 1).text == "2"

    srt = client.get(f"/api/meetings/{meeting['id']}/export", params={"format": "srt"}).text
    assert srt.startswith("1\n00:00:00,000 --> 00:00:02,500\nBom dia a todos.")
    txt = client.get(f"/api/meetings/{meeting['id']}/export", params={"format": "txt"}).text
    assert txt == "Bom dia a todos.\nVamos começar.\n"


def test_context_files_reach_the_prompt(settings: Settings, transcriber: FakeTranscriber, gpu: FakeGpu) -> None:
    settings.context_dir.mkdir(parents=True)
    (settings.context_dir / "glossario.md").write_text("VGV = valor geral de vendas", encoding="utf-8")
    chat = FakeChat()
    with TestClient(create_app(settings, transcriber=transcriber, gpu=gpu, chat=chat)) as client:
        meeting = transcribed_meeting(client)
        job = client.post(f"/api/meetings/{meeting['id']}/minutes/generate").json()
        wait_for(client, job["id"], "done")
    assert "VGV = valor geral de vendas" in chat.calls[-1][2]
